from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
PARENT_ROOT = PROJECT_ROOT.parent
REPORTS_DIR = PROJECT_ROOT / "data" / "reports"
OUTPUT = REPORTS_DIR / "raw_file_inventory.csv"

DATA_EXTENSIONS = {
    ".csv",
    ".xlsx",
    ".xls",
    ".json",
    ".parquet",
    ".html",
    ".htm",
    ".txt",
    ".yaml",
    ".yml",
}


def source_id_from_path(path: Path) -> str:
    rel = path.relative_to(PROJECT_ROOT) if path.is_relative_to(PROJECT_ROOT) else path.name
    parts = Path(rel).parts if isinstance(rel, Path) else (str(rel),)
    if "raw" in parts:
        idx = parts.index("raw")
        if len(parts) > idx + 1:
            return parts[idx + 1]
    if "bronze" in parts:
        idx = parts.index("bronze")
        if len(parts) > idx + 1:
            return parts[idx + 1]
    if path.name == "south_korea_ai_index_dashboard_data.csv":
        return "south_korea_ai_index_dashboard_data"
    if path.parent.name in {"sources", "data", "docs"}:
        return path.stem
    return path.parent.name or path.stem


def count_csv(path: Path) -> tuple[str, str]:
    try:
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            reader = csv.reader(handle)
            header = next(reader, [])
            rows = sum(1 for _ in reader)
        return str(rows), "|".join(header)
    except Exception as exc:
        return "", f"READ_ERROR: {exc}"


def inspect_tabular(path: Path) -> tuple[str, str, str]:
    ext = path.suffix.lower()
    if ext == ".csv":
        rows, columns = count_csv(path)
        return rows, columns, ""
    if ext in {".xlsx", ".xls"}:
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
            details = []
            total_rows = 0
            for ws in wb.worksheets:
                rows = max(ws.max_row - 1, 0)
                total_rows += rows
                headers = [str(c.value).strip() for c in next(ws.iter_rows(min_row=1, max_row=1), []) if c.value is not None]
                details.append(f"{ws.title}: {'|'.join(headers)}")
            wb.close()
            return str(total_rows), " ; ".join(details), ""
        except Exception as exc:
            return "", "", f"READ_ERROR: {exc}"
    if ext == ".parquet":
        try:
            df = pd.read_parquet(path)
            return str(len(df)), "|".join(map(str, df.columns)), ""
        except Exception as exc:
            return "", "", f"READ_ERROR: {exc}"
    if ext == ".json":
        try:
            obj: Any = json.loads(path.read_text(encoding="utf-8", errors="replace"))
            if isinstance(obj, list):
                cols = sorted({k for row in obj if isinstance(row, dict) for k in row.keys()})
                return str(len(obj)), "|".join(cols), ""
            if isinstance(obj, dict):
                return "1", "|".join(obj.keys()), ""
            return "", "", ""
        except Exception as exc:
            return "", "", f"READ_ERROR: {exc}"
    return "", "", ""


def iter_input_files() -> list[Path]:
    roots = [
        PROJECT_ROOT / "data",
        PROJECT_ROOT / "sources",
        PROJECT_ROOT / "docs",
    ]
    files: list[Path] = []
    for root in roots:
        if root.exists():
            files.extend(path for path in root.rglob("*") if path.is_file() and path.suffix.lower() in DATA_EXTENSIONS)
    external = PARENT_ROOT / "south_korea_ai_index_dashboard_data.csv"
    if external.exists():
        files.append(external)
    return sorted(set(files), key=lambda p: str(p).lower())


def main() -> None:
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    rows = []
    for path in iter_input_files():
        stat = path.stat()
        row_count, columns, note = inspect_tabular(path)
        try:
            display_path = str(path.relative_to(PROJECT_ROOT))
        except ValueError:
            display_path = str(path)
        rows.append(
            {
                "file_path": display_path,
                "absolute_file_path": str(path),
                "file_type": path.suffix.lower().lstrip(".") or "no_extension",
                "file_size": stat.st_size,
                "modified_date": stat.st_mtime,
                "rows_if_tabular": row_count,
                "columns_if_tabular": columns,
                "detected_source_id": source_id_from_path(path),
                "notes": note,
            }
        )
    pd.DataFrame(rows).to_csv(OUTPUT, index=False, encoding="utf-8-sig")
    print(f"Input files found: {len(rows)}")
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
