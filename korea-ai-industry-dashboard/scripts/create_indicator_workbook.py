from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from kai_dash.config import COLLECTION_STATUS_COLUMNS, MANUAL_REVIEW_COLUMNS
from kai_dash.data_io import indicators_df, read_collection_status, read_manual_review, sources_df
from kai_dash.paths import DOCS_DIR, ensure_dirs


DASHBOARD_PAGES = [
    {
        "page_name": "Home",
        "purpose": "Research dashboard entry point with coverage and key collected values.",
        "indicators_used": "All collected gold observations",
        "charts_planned": "Coverage KPIs, latest observed values, collection status",
        "status": "implemented",
    },
    {
        "page_name": "01_Overview",
        "purpose": "Cross-domain first-pass view of Korea AI industry signals.",
        "indicators_used": "Priority A/B collected indicators",
        "charts_planned": "Latest-value comparison bars and observation table",
        "status": "implemented",
    },
    {
        "page_name": "02_Compute_and_Data_Centers",
        "purpose": "Compute, public GPU targets, data-center envelope, and known cluster indicators.",
        "indicators_used": "compute category",
        "charts_planned": "GPU targets, KISTI capacity, TOP500, known clusters",
        "status": "implemented",
    },
    {
        "page_name": "03_Hardware_and_Semiconductors",
        "purpose": "Semiconductor trade and Korean hardware company context.",
        "indicators_used": "semiconductors category",
        "charts_planned": "Semiconductor trade and filing-derived metrics as available",
        "status": "implemented",
    },
    {
        "page_name": "04_Models_Research_Funding",
        "purpose": "Frontier model, research publication, patent, and funding indicators.",
        "indicators_used": "models, research, funding, startups",
        "charts_planned": "Epoch model counts, OpenAlex publications, Stanford AI Index values",
        "status": "implemented",
    },
    {
        "page_name": "05_Adoption_and_Industry",
        "purpose": "Industrial automation and worker/firm adoption.",
        "indicators_used": "adoption, worker_adoption, productivity, firm_adoption",
        "charts_planned": "Robot density/installations and BOK worker adoption metrics",
        "status": "implemented",
    },
    {
        "page_name": "06_Talent_Public_Defense",
        "purpose": "Talent, public sector, governance context, and defense-public proxies.",
        "indicators_used": "talent, public_sector, governance, defense",
        "charts_planned": "Collected public budget, defense budget proxy, talent and governance as available",
        "status": "implemented",
    },
    {
        "page_name": "07_Source_Explorer",
        "purpose": "Inspect source registry and downloaded-source status.",
        "indicators_used": "source registry",
        "charts_planned": "Source table and filters",
        "status": "implemented",
    },
    {
        "page_name": "08_Collection_Report",
        "purpose": "Review collection status, manual queue, and coverage report.",
        "indicators_used": "collection_status, manual_review_queue, data_coverage",
        "charts_planned": "Status tables",
        "status": "implemented",
    },
]


README_ROWS = [
    ("Project goal", "Build a v0.1 research dashboard with straightforward public indicators for South Korea's AI industry. It is not a complete country profile."),
    ("Collection approach", "The collector downloads public source material into data/raw/{source_id}/{YYYY-MM-DD}/ and writes metadata JSON for auditability."),
    ("Pipeline layers", "Raw stores unedited source material. Bronze stores lightly parsed source tables. Silver stores normalized observations. Gold stores dashboard-ready Parquet and DuckDB files."),
    ("Confidence labels", "High means official or primary source with clear value. Medium_high means reputable dataset/report with methodology caveats. Medium means estimates, surveys, market reports, or manual classification."),
    ("Indicator types", "Observed metrics, official targets, estimates, company claims, index metrics, survey metrics, manual classifications, contextual facts, and missing/manual-review items are separated."),
    ("Interpretation warning", "Official targets, estimates, and company claims are not the same as observed deployed capacity or audited operating metrics."),
]


def add_dataframe_sheet(wb: Workbook, name: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(name)
    if df.empty:
        df = pd.DataFrame(columns=list(df.columns))
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([excel_safe(value) for value in row])
    style_sheet(ws)
    if ws.max_column and ws.max_row:
        ref = f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, 1)}"
        table = Table(displayName=name.replace(" ", "_")[:30], ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)


def excel_safe(value):
    if isinstance(value, list):
        return ", ".join(map(str, value))
    if isinstance(value, dict):
        return str(value)
    if pd.isna(value):
        return ""
    return value


def style_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        values = [str(cell.value) if cell.value is not None else "" for cell in column_cells]
        width = min(max(len(v) for v in values) + 2, 55)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = max(width, 12)


def main() -> None:
    ensure_dirs()
    indicators = indicators_df()
    sources = sources_df()
    collection = read_collection_status()
    if collection.empty:
        collection = pd.DataFrame(columns=COLLECTION_STATUS_COLUMNS)
    manual = read_manual_review()
    if manual.empty:
        manual = pd.DataFrame(columns=MANUAL_REVIEW_COLUMNS)
    pages = pd.DataFrame(DASHBOARD_PAGES)

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme.append(["Topic", "Notes"])
    for row in README_ROWS:
        readme.append(list(row))
    style_sheet(readme)
    ref = f"A1:B{readme.max_row}"
    table = Table(displayName="README_Table", ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    readme.add_table(table)

    add_dataframe_sheet(wb, "Indicator_Catalog", indicators)
    add_dataframe_sheet(wb, "Source_Registry", sources)
    add_dataframe_sheet(wb, "Collection_Status", collection)
    add_dataframe_sheet(wb, "Manual_Review_Queue", manual)
    add_dataframe_sheet(wb, "Dashboard_Pages", pages)

    output = DOCS_DIR / "korea_ai_v01_indicator_plan.xlsx"
    wb.save(output)
    print(output)


if __name__ == "__main__":
    main()
