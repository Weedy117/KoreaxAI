from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
CONSOLIDATED_DIR = PROJECT_ROOT / "data" / "consolidated"
REPORTS_DIR = PROJECT_ROOT / "data" / "reports"
MANUAL_REVIEW_DIR = PROJECT_ROOT / "data" / "manual_review"
DOCS_DIR = PROJECT_ROOT / "docs"
OUTPUT = DOCS_DIR / "korea_ai_consolidated_data_workbook.xlsx"


def read_csv(name: str, folder: Path = CONSOLIDATED_DIR) -> pd.DataFrame:
    path = folder / name
    if not path.exists():
        return pd.DataFrame()
    return pd.read_csv(path, dtype=str, keep_default_na=False)


def make_comparison_wide(canonical: pd.DataFrame, coverage: pd.DataFrame) -> pd.DataFrame:
    if canonical.empty:
        return pd.DataFrame()
    usable = canonical[canonical["country_iso3"].isin(["KOR", "USA", "CHN"])].copy()
    usable = usable[usable["value_standardized"].astype(str) != ""]
    if usable.empty:
        return pd.DataFrame()
    usable["value_for_display"] = usable["value_standardized"].where(
        usable["value_standardized"].astype(str) != "",
        usable["value_raw"],
    )
    idx_cols = ["indicator_id", "indicator_name", "category", "year", "unit_standardized"]
    wide = usable.pivot_table(
        index=idx_cols,
        columns="country_iso3",
        values="value_for_display",
        aggfunc="first",
    ).reset_index()
    wide.columns.name = None
    wide = wide.rename(
        columns={
            "unit_standardized": "unit",
            "KOR": "kor_value",
            "USA": "usa_value",
            "CHN": "chn_value",
        }
    )
    for col in ["kor_value", "usa_value", "chn_value"]:
        if col not in wide.columns:
            wide[col] = ""
    coverage_small = coverage[["indicator_id", "comparison_status", "notes"]].rename(columns={"notes": "source_note"})
    return wide.merge(coverage_small, on="indicator_id", how="left")[
        [
            "indicator_id",
            "indicator_name",
            "category",
            "year",
            "kor_value",
            "usa_value",
            "chn_value",
            "unit",
            "comparison_status",
            "source_note",
        ]
    ].sort_values(["category", "indicator_name", "year"])


def readme_df() -> pd.DataFrame:
    sections = [
        ("Purpose", "A source-governed v0.1 dataset for a graduate dashboard on AI development and deployment in South Korea."),
        ("Data pipeline", "Inventory local files, map structured observations into a long master schema, normalize countries/periods/units, deduplicate, select canonical rows, then export CSV, Excel, Parquet, and DuckDB outputs."),
        ("Indicator types", "Observed metrics, official targets, estimates, company claims, index metrics, survey metrics, manual classifications, contextual facts, and missing placeholders are kept distinct."),
        ("Confidence tiers", "high, medium_high, medium, medium_low, low, and manual_review."),
        ("Canonical value selection", "Canonical rows are selected by source hierarchy, confidence tier, indicator type, and recency. Alternatives remain in master_observations_all.csv."),
        ("Deduplication rules", "Exact duplicate rows are removed and logged. Competing duplicate groups are preserved and queued for review."),
        ("Normalization rules", "Country names map to ISO3; annual, quarterly, monthly, and event periods are standardized; values preserve raw units and only safe conversions are applied."),
        ("Comparison coverage rules", "Coverage is evaluated for KOR, USA, and CHN based on canonical observations; no comparison values are invented."),
        ("Known limitations", "Some raw HTML and text sources are inventoried but not table-extracted in this pass; currency conversion is intentionally not performed without a sourced exchange rate."),
    ]
    return pd.DataFrame(sections, columns=["Section", "Description"])


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    title_fill = PatternFill("solid", fgColor="1F4E78")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="000000")
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col_idx, column_cells in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in list(column_cells)[:200]:
                max_len = max(max_len, len(str(cell.value or "")))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 48)
        for row in ws.iter_rows():
            for cell in row:
                if cell.column > 8 or "note" in str(ws.cell(row=1, column=cell.column).value).lower():
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
    readme = wb["README"]
    readme.insert_rows(1, 3)
    readme["A1"] = "South Korea AI Industry Dashboard — Consolidated Dataset"
    readme["A2"] = "A source-governed v0.1 dataset for tracking AI development, deployment, and strategic industry markers in South Korea, with US and China comparisons where available."
    readme["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    readme["A1"].fill = title_fill
    readme["A2"].font = Font(italic=True, size=11)
    readme.merge_cells("A1:B1")
    readme.merge_cells("A2:B2")
    readme.freeze_panes = "A5"
    wb.save(path)


def main() -> None:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    canonical = read_csv("master_observations_canonical.csv")
    timeseries = read_csv("master_observations_timeseries.csv")
    indicators = read_csv("master_indicators.csv")
    sources = read_csv("master_sources.csv")
    duplicates = read_csv("duplicates_removed.csv")
    log = read_csv("normalization_log.csv")
    review = read_csv("data_quality_review_queue.csv", MANUAL_REVIEW_DIR)
    coverage = read_csv("comparison_coverage.csv")
    inventory = read_csv("raw_file_inventory.csv", REPORTS_DIR)

    key_cols = [
        "category",
        "indicator_name",
        "year",
        "value_standardized",
        "unit_standardized",
        "indicator_type",
        "confidence_tier",
        "source_name",
        "notes",
    ]
    korea_key = canonical[canonical.get("country_iso3", "") == "KOR"][key_cols].copy() if not canonical.empty else pd.DataFrame(columns=key_cols)
    comparison = make_comparison_wide(canonical, coverage)

    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        readme_df().to_excel(writer, sheet_name="README", index=False)
        canonical.to_excel(writer, sheet_name="Dashboard_Master", index=False)
        timeseries.to_excel(writer, sheet_name="Time_Series_Data", index=False)
        korea_key.to_excel(writer, sheet_name="Korea_Key_Indicators", index=False)
        comparison.to_excel(writer, sheet_name="Korea_US_China_Comparison", index=False)
        indicators.to_excel(writer, sheet_name="Indicator_Dictionary", index=False)
        sources.to_excel(writer, sheet_name="Source_Index", index=False)
        duplicates.to_excel(writer, sheet_name="Duplicates_Removed", index=False)
        log.to_excel(writer, sheet_name="Normalization_Log", index=False)
        review.to_excel(writer, sheet_name="Manual_Review_Queue", index=False)
        coverage.to_excel(writer, sheet_name="Data_Coverage", index=False)
        inventory.to_excel(writer, sheet_name="Raw_File_Inventory", index=False)

    style_workbook(OUTPUT)
    print(f"Wrote workbook: {OUTPUT}")


if __name__ == "__main__":
    main()
