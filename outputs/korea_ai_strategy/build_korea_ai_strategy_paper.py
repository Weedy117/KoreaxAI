from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path(r"C:\Users\ASUS\Desktop\SAIS Playground\ROKxAI\outputs\korea_ai_strategy")
FIG_DIR = OUT_DIR / "figures"
DOCX_PATH = OUT_DIR / "korea_ai_strategy_paper.docx"


NOTES = {
    "ntsf": "National Technology Strategy Framework, Version 2, class discussion document, 2026.",
    "syllabus": "Johns Hopkins SAIS, Technology, Power, and Statecraft syllabus, Spring 2026.",
    "stanford": "Stanford Institute for Human Centered Artificial Intelligence, The 2026 AI Index Report, 2026.",
    "dashboard": "ROKxAI dashboard, consolidated public source dataset, collected May 3, 2026.",
    "strategy_2019": "Ministry of Science and ICT, National Strategy for Artificial Intelligence, December 17, 2019.",
    "digital_strategy": "Ministry of Science and ICT, National Digital Strategy announcement, September 28, 2022.",
    "ai_basic": "Ministry of Science and ICT, press release on the Basic Act on the Development of Artificial Intelligence and the Establishment of a Trustworthy Foundation entering into force, January 22, 2026.",
    "ai_committee": "Ministry of Science and ICT, press release on establishment of the National AI Strategy Committee, September 16, 2025.",
    "budget": "Ministry of Science and ICT, press release on the KRW 1.9067 trillion supplementary AI budget, May 19, 2025.",
    "gpu_procurement": "Ministry of Science and ICT, press release on GPU Procurement Project participants, July 2025.",
    "work_plan": "Ministry of Science and ICT, 2026 Work Plan, December 12, 2025.",
    "models": "Ministry of Science and ICT, press release on the Sovereign AI Foundation Model project, August 18, 2025.",
    "physical_ai": "Ministry of Science and ICT, press release on the Physical AI Global Alliance, September 29, 2025.",
    "nvidia": "NVIDIA, press release on Korea AI infrastructure with the Korean government and industrial firms, October 30, 2025.",
    "trendforce": "TrendForce, press release on second quarter DRAM revenue and supplier shares, September 2, 2025.",
    "oecd": "OECD and Korea Labor Institute, Artificial Intelligence and the Labour Market in Korea, OECD Publishing, 2025.",
    "defense": "Yonhap News Agency, report on the launch of the Defense AI Center, April 1, 2024.",
    "pipc": "Personal Information Protection Commission, press release on the prior adequacy review of Kakao Kanana, March 13, 2025.",
    "pipa": "Personal Information Protection Commission, Personal Information Protection Act materials on pseudonymized information, accessed May 5, 2026.",
    "horowitz": "Michael C. Horowitz, Artificial Intelligence, International Competition, and the Balance of Power, Texas National Security Review 1, no. 3, May 2018.",
    "ding": "Jeffrey Ding, Technology and the Rise of Great Powers: How Diffusion Shapes Economic Competition, Princeton University Press, 2024.",
    "rubber": "American Chemical Society, United States Synthetic Rubber Program, 1939 to 1945, National Historic Chemical Landmark booklet, 1998.",
    "wellerstein": "Alex Wellerstein, Manhattan Project, Encyclopedia of the History of Science, 2019.",
    "ohno": "Kenichi Ohno, Meiji Japan: Progressive Learning of Western Technology, in How Nations Learn: Technological Learning, Industrial Policy, and Catch Up, edited by Arkebe Oqubay and Kenichi Ohno, Oxford University Press, 2019.",
    "bush": "Vannevar Bush, Science, The Endless Frontier, report to the President, 1945.",
    "nscai": "National Security Commission on Artificial Intelligence, Final Report, 2021.",
    "singapore": "Government of Singapore, National Artificial Intelligence Strategy 2.0, 2023, and Smart Nation 2.0, 2024.",
    "taiwan": "Taiwan National Development Council, Ten AI Initiatives Promotion Plan, 2025 to 2028, approved January 28, 2026.",
    "israel": "Israel Innovation Authority, The State of High Tech, Annual Report 2025.",
}


ORDERED_NOTE_KEYS: list[str] = []


def note_number(key: str) -> int:
    if key not in ORDERED_NOTE_KEYS:
        ORDERED_NOTE_KEYS.append(key)
    return ORDERED_NOTE_KEYS.index(key) + 1


def reject_dashes(text: str, label: str) -> None:
    bad = [ch for ch in text if ch in "-–—"]
    if bad:
        raise ValueError(f"Dash found in {label}: {text}")


def add_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    reject_dashes(text, "table cell")
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(8.5)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_citation(paragraph, key: str) -> None:
    n = note_number(key)
    r = paragraph.add_run(str(n))
    r.font.superscript = True
    r.font.size = Pt(8)


def add_citations(paragraph, keys: list[str]) -> None:
    if not keys:
        return
    refs = [str(note_number(key)) for key in keys]
    r = paragraph.add_run(",".join(refs))
    r.font.superscript = True
    r.font.size = Pt(8)


def add_para(doc: Document, text: str, cites: list[str] | None = None, style: str | None = None):
    reject_dashes(text, "paragraph")
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.08
    p.add_run(text)
    add_citations(p, cites or [])
    return p


def add_heading(doc: Document, text: str, level: int = 1):
    reject_dashes(text, "heading")
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_numbered(doc: Document, items: list[tuple[str, list[str]]]):
    for text, cites in items:
        if text.startswith("Industrial adoption") or text.startswith("Create a national talent"):
            doc.add_page_break()
        p = add_para(doc, text, cites, style="List Number")
        p.paragraph_format.left_indent = Inches(0.25)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], title: str | None = None):
    if title:
        cap = add_para(doc, title)
        cap.runs[0].bold = True
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color="FFFFFF")
        add_shading(table.rows[0].cells[i], "1F4E5F")
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
    doc.add_paragraph()
    return table


def add_figure(doc: Document, path: Path, caption: str):
    reject_dashes(caption, "figure caption")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(path), width=Inches(6.4))
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.runs[0].italic = True
    c.runs[0].font.size = Pt(8.5)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.2)
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 95)
    styles["Heading 2"].font.name = "Aptos"
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].font.color.rgb = RGBColor(31, 78, 95)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(22)
    styles["Title"].font.bold = True


def build_doc() -> Document:
    doc = Document()
    style_document(doc)

    title = "Korea's AI Strategy as Statecraft"
    subtitle = "Industrial Capacity, Trusted Data, and Middle Power Leverage"
    reject_dashes(title, "title")
    reject_dashes(subtitle, "subtitle")
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    q = doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    qr = q.add_run(subtitle)
    qr.font.size = Pt(14)
    qr.font.color.rgb = RGBColor(80, 80, 80)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Policy research paper for Technology, Power, and Statecraft").italic = True
    doc.add_paragraph()

    add_heading(doc, "Executive Summary")
    add_para(
        doc,
        "South Korea's AI strategy should be judged by whether it can convert a compact set of industrial and institutional advantages into durable strategic influence. Frontier model parity with the United States and China is an insufficient benchmark because it treats AI power as a contest over model counts rather than as a struggle over control points, adoption capacity, data legitimacy, and alliance value. Korea's more plausible route is to connect high bandwidth memory, advanced manufacturing, dense robotics deployment, Korean language and domain data, privacy credible regulation, and defense adoption into an integrated AI industrial system.",
        ["stanford", "dashboard"],
    )
    add_para(
        doc,
        "The empirical profile gives the strategy a physical center of gravity. The Stanford AI Index identifies Korea as the world leader in AI patents per capita, and the project dashboard records Korea's exceptional manufacturing robot density relative to the comparison set. Samsung Electronics and SK hynix occupy central positions in global DRAM and high bandwidth memory supply. These assets locate Korea inside the material infrastructure of AI rather than only inside its software layer. They also make the country a credible test bed for AI in manufacturing, defense, mobility, health, and public services, where deployment discipline can create advantages that model benchmarks do not capture.",
        ["dashboard", "trendforce"],
    )
    add_para(
        doc,
        "The central weakness is institutional coupling. The Ministry of Science and ICT has moved quickly through the AI Basic Act, the National AI Strategy Committee, a large AI budget, public GPU procurement, and five sovereign foundation model teams, but these instruments do not yet amount to a mission architecture that binds models, compute, data, standards, and procurement to priority national problems. A model centered strategy can produce capable Korean systems while leaving strategic rents in foreign cloud platforms, foreign GPU ecosystems, and foreign application markets. Sovereignty in AI therefore depends on control over the full stack of inputs, channels, legal permissions, assurance practices, and demand.",
        ["ai_basic", "ai_committee", "budget", "models"],
    )
    add_para(
        doc,
        "The recommendations follow from that diagnosis. Korea should treat AI strategy as stack statecraft by using memory as alliance currency, building a national compute and data commons, turning K AI models into a shared deployment layer, making manufacturing and defense anchor customers, converting the Personal Information Protection Act regime into trusted data infrastructure, exporting implementation practice across the Indo Pacific, and measuring success through industrial diffusion and institutional leverage rather than benchmark publicity.",
        ["ntsf", "nscai"],
    )

    add_heading(doc, "Argument and Framework")
    add_para(
        doc,
        "AI becomes a technology of statecraft for Korea when industrial bottlenecks are connected to diplomatic, economic, and military objectives. The National Technology Strategy Framework clarifies this connection by asking why the state commits to a technology, how it organizes science and industry, and how capacity is converted into power. Korea's answer is shaped by five pressures: the North Korean threat, demographic contraction, dependence on external frontier model and GPU ecosystems, the productivity needs of an advanced manufacturing economy, and the opportunity to turn memory chip dominance into geopolitical leverage. These pressures make AI a means of prosperity, deterrence, social resilience, and alliance bargaining.",
        ["ntsf", "oecd", "defense"],
    )
    add_para(
        doc,
        "The historical relevance of Korea's developmental state lies less in a direct return to command planning than in the persistence of institutional habits after democratization: presidential agenda setting, ministry coordination, framework legislation, targeted finance, public private consultation, and bargaining with chaebol and platform firms. AI extends this pattern into a harder domain because the state must coordinate compute infrastructure, privacy law, safety obligations, procurement demand, semiconductor supply, and alliance commitments at the same time. The AI Basic Act is therefore significant not as a self executing solution, but as a framework statute whose strategic value depends on subordinate regulations, guidance, enforcement practice, and administrative capacity.",
        ["ai_basic", "pipa"],
    )
    add_para(
        doc,
        "Korea's operating model is hybrid, which gives it flexibility and creates coordination risk. It relies on national champions in semiconductors, telecommunications, internet platforms, automobiles, and electronics. It uses a consortium logic for sovereign foundation models and public compute, a state project logic for defense AI and public services, and a commercial dual use logic because most deployable AI capability will come from firms that need profits, customers, exports, and developer ecosystems. The policy problem is therefore institutional design: Korea must make these models mutually reinforcing rather than allowing each instrument to produce separate constituencies, standards, and procurement channels.",
        ["ntsf", "models", "nvidia"],
    )

    add_figure(
        doc,
        FIG_DIR / "figure_1_patents_robot_density.png",
        "Figure 1. Korea combines high AI patent density with exceptional manufacturing automation.",
    )

    add_heading(doc, "Strategic Landscape")
    add_heading(doc, "State of Play: From Digital Catch Up to Sovereign AI", 2)
    add_para(
        doc,
        "Korea's strategic landscape is best read as a transition from digital modernization to AI stack management. Earlier policy sought to make the state and the economy more digital; the current agenda treats AI as infrastructure for industrial competition, alliance bargaining, demographic resilience, and military adaptation. The transition remains incomplete because MSIT, the Personal Information Protection Commission, the Ministry of National Defense, industrial ministries, the presidential AI strategy body, and the National Assembly govern different pieces of the AI system. Korea is not positioned to match the United States or China in general frontier scale. Its more credible path is to become indispensable in memory, trusted deployment, manufacturing AI, and allied supply assurance.",
        ["ntsf", "stanford", "dashboard", "ai_committee"],
    )
    add_para(
        doc,
        "The policy trajectory shows four phases rather than a single acceleration story. The 2019 National Strategy treated AI as an economy wide modernization agenda organized around AI ecosystem building, AI utilization, and people centered AI. The Digital New Deal and Data Dam phase converted that ambition into data infrastructure and public digital investment. The 2022 National Digital Strategy broadened the agenda to AI chips, data, platform government, talent, and global digital norms. Since the National Assembly passed the AI Basic Act in December 2024 and the presidential AI strategy body was relaunched in 2025, the center of gravity has shifted toward strategic AI through legal authority, GPU procurement, sovereign foundation models, the KRW 1.9067 trillion supplementary AI budget, and the 2026 Work Plan. The result is a more explicit theory of national power in which AI is treated as an infrastructure layer for industry, science, security, and public administration.",
        ["strategy_2019", "digital_strategy", "ai_basic", "ai_committee", "budget", "models", "work_plan"],
    )
    add_figure(
        doc,
        FIG_DIR / "figure_4_policy_timeline.png",
        "Figure 2. Korea's AI policy evolution shows a shift from digital modernization to stack sovereignty.",
    )
    add_heading(doc, "Geopolitical Stack Position", 2)
    add_para(
        doc,
        "The global AI hierarchy is organized through a stack in which states and firms hold different sources of leverage. The United States dominates frontier model companies, cloud platforms, GPUs, and software ecosystems. China supplies scale, state backed deployment, and an alternative platform system. Taiwan is central to advanced foundry capacity. Korea's position is narrower but strategically valuable: it contains world class memory firms, dense manufacturing capability, major digital platforms, and some of the most automation intensive industrial environments in the world. This position creates leverage only if Seoul converts component importance into negotiated access to compute, design tools, export markets, and allied research networks.",
        ["stanford", "dashboard", "trendforce"],
    )
    add_para(
        doc,
        "High bandwidth memory gives Korea a concrete bargaining instrument in the AI race, although that instrument is conditional rather than automatic. TrendForce reported that SK hynix led the DRAM market by revenue in the second quarter of 2025 with 38.7 percent share, while Samsung ranked second with 32.7 percent. Frontier AI systems depend on memory bandwidth, packaging reliability, power efficiency, and stable supply as much as they depend on model design. Korea can therefore use memory and advanced DRAM as alliance assets, linking supply assurance to reciprocal access to GPUs, cloud capacity, trusted data flows, and standards forums. The danger is that memory remains a profitable input while the strategic rents from models, cloud platforms, and developer ecosystems accrue elsewhere.",
        ["trendforce", "nvidia"],
    )
    add_heading(doc, "Domestic Demand and Physical AI", 2)
    add_para(
        doc,
        "Korea's domestic demand for AI is shaped less by consumer novelty than by structural pressure in manufacturing, logistics, public services, health care, and defense. The OECD and Korea Labor Institute warn that demographic aging could push the old age dependency ratio above 75 percent by 2060 and reduce the working age population by up to 46 percent between 2023 and 2060. In that context, AI is a productivity instrument for an advanced economy whose labor supply is tightening. This gives Korea a stronger rationale for physical AI than many peer economies: models must be embedded in robots, machine tools, shipyards, vehicles, hospitals, grid systems, and public administration rather than remaining primarily as conversational software.",
        ["oecd", "work_plan", "physical_ai"],
    )
    add_para(
        doc,
        "The dashboard shows why this adoption layer is both an advantage and a policy problem. Korea recorded robot density of 1,220 robots per 10,000 manufacturing workers in 2024, far above the United States and China in the comparison set, but annual industrial robot installations have been broadly flat at about 31,000 units since 2020. The country also shows rising AI engineering skills diffusion, reaching 16.8 percent in 2025, and worker use of generative AI at work above half of surveyed workers. These figures suggest a strong base for industrial absorption, but they also imply that the next margin of competitiveness will come from software integration, worker redesign, safety assurance, and sector specific deployment rather than additional automation alone.",
        ["dashboard"],
    )
    doc.add_page_break()
    add_heading(doc, "Security and Military Adoption", 2)
    add_para(
        doc,
        "Korea's security environment pushes AI toward force multiplication under demographic and geographic constraint. The North Korean missile, artillery, cyber, and unmanned systems threat demands faster sensing, decision support, logistics, and cyber defense, while manpower decline makes a larger military less plausible over time. The Defense AI Center, launched at the Agency for Defense Development in Daejeon in April 2024, gives institutional form to the Defense Innovation 4.0 agenda by linking civilian and military personnel, manned and unmanned teaming, battlefield awareness, and private sector technology transfer. This is not only a weapons question. It is an organizational question about how to integrate AI into command, maintenance, training, electronic warfare, intelligence analysis, and military exercises while preserving human responsibility for escalation sensitive decisions.",
        ["defense", "horowitz"],
    )
    add_heading(doc, "Data Sovereignty and Trust", 2)
    add_para(
        doc,
        "Sovereign AI requires legitimate data mobilization rather than national ownership alone. Korea's large platforms hold valuable social, search, commerce, communication, and payment data, but democratic legitimacy, competition concerns, and privacy law limit how those assets can be used. The AI Basic Act gives the state a promotion and trust framework, including national governance, support for research and data centers, transparency rules for generative AI, and high impact AI obligations in areas such as energy, drinking water, health care, nuclear facilities, criminal investigations, recruitment, credit, transport, public services, and education. Those rules will become industrial infrastructure only if they work with the Personal Information Protection Act rather than around it.",
        ["ai_basic", "pipa"],
    )
    add_para(
        doc,
        "The Personal Information Protection Act already contains a pathway for data use through pseudonymized information for purposes such as statistics, scientific research, and public interest archiving, while the Personal Information Protection Commission remains the central regulator that gives those pathways practical shape. The commission's prior adequacy review of Kakao Kanana illustrates the mechanism Korea must scale: the regulator permitted innovation while requiring separate storage, encryption, constraints on OpenAI use, consent rules for training, and continuing oversight. Data governance is therefore not simply a brake on industrial policy. Properly designed, it can create trusted data spaces that give firms lawful access to high value data while preserving public confidence. Korea's comparative advantage will weaken if privacy remains a case by case compliance burden; it will strengthen if the state turns privacy review into repeatable institutions for safe data sharing.",
        ["pipa", "pipc"],
    )

    add_heading(doc, "How Korea Is Building Capacity")
    add_para(
        doc,
        "Korea is constructing AI capacity through bottleneck governance rather than through a single technology program. State instruments supply legal authority, finance, compute, procurement, and coordination. Science supplies patents, researchers, benchmarks, domain knowledge, and absorptive capacity. Industry supplies memory, fabs, platforms, robotics, customers, and export channels. The central policy challenge is whether these layers will combine into a national learning system or remain a portfolio of impressive but weakly connected initiatives.",
        ["ntsf", "ai_basic"],
    )
    add_heading(doc, "Policy Architecture and Strategic Sequencing", 2)
    add_para(
        doc,
        "The state pillar has become more operational since 2025 because Korea now combines legal authority with budgetary allocation, subordinate rulemaking, and named implementation vehicles. The AI Basic Act creates the statutory foundation, while its Enforcement Decree, transparency guidelines, grace period, and support desk illustrate Korea's regulatory culture of framework legislation followed by administrative interpretation. The supplementary budget concentrates resources on compute, foundation models, talent, and domestic AI semiconductors. The 2026 Work Plan raises the AI budget from KRW 3.0 trillion to KRW 9.9 trillion and organizes policy around the AI Highway, K AI, AI co scientists, regional AI transformation, and public services. Korea has moved from capability aspiration to bottleneck management.",
        ["ai_basic", "budget", "work_plan"],
    )
    add_para(
        doc,
        "Bottleneck management gives the Korean state leverage only if allocation is connected to performance discipline. Public GPUs should not become a general subsidy for incumbents, and sovereign models should not become symbolic national artifacts. The stronger design is to tie compute, data, and talent support to public benchmarks, industrial deployment targets, open interfaces, security evaluation, privacy compliance, and measurable diffusion into small firms, regional clusters, public services, and defense users. This is where the NTSF state pillar must discipline the industry pillar rather than merely finance it.",
        ["ntsf", "budget", "models", "pipa"],
    )
    add_para(
        doc,
        "AI capacity crosses statutory boundaries. MSIT leads the industrial and digital agenda, but the Personal Information Protection Commission, the Ministry of National Defense, industrial ministries, and the National Assembly control data legitimacy, military demand, manufacturing adoption, energy intensive facilities, and law. The National AI Strategy Committee adds value only if it turns those authorities into shared milestones and visible tradeoffs.",
        ["ai_committee", "ai_basic", "pipa", "defense"],
    )
    add_heading(doc, "Compute and Model Layer", 2)
    add_para(
        doc,
        "Compute is Korea's most urgent capacity constraint because model development, scientific AI, defense testing, and startup scaling all depend on access to clustered GPUs. The 2025 supplementary budget initially targeted 10,000 advanced GPUs, and the procurement project later selected Naver Cloud, NHN Cloud, and Kakao to secure 13,000 GPUs, including B200 and H200 systems. MSIT's 2026 plan then set a cumulative target of 37,000 GPUs by 2026 through government procurement and Supercomputer Number 6. NVIDIA's October 2025 announcement is larger still, with the Korean government and major firms adding more than 260,000 GPUs across sovereign clouds and AI factories. The scale is significant, but it also reveals dependence on foreign accelerator supply, CUDA centered software, data center power, and allied export control decisions.",
        ["budget", "gpu_procurement", "work_plan", "nvidia"],
    )
    add_para(
        doc,
        "The Sovereign AI Foundation Model project uses a managed tournament to convert scarce inputs into model capability. Naver Cloud, Upstage, SK Telecom, NC AI, and LG AI Research were selected after review, with support in GPUs, data, and talent. Their goals cover omni models, public AI services, industrial models, agents, manufacturing, defense, health, finance, robotics, and public services. The mechanism is a hybrid of the NTSF consortium model and national champion logic: the state pools scarce resources and defines national objectives while firm rivalry supplies technical discipline. The tournament will create public value only if models remain interoperable enough for procurement, safety review, sector adapters, and startup access.",
        ["models", "ntsf"],
    )
    add_figure(
        doc,
        FIG_DIR / "figure_2_compute_model_base.png",
        "Figure 3. Korea is increasing compute and model capacity, but the base remains narrow compared with superpower scale.",
    )
    add_heading(doc, "Innovation Layer: Patents, Research, and Models", 2)
    add_para(
        doc,
        "Korea's innovation layer is dense, technically credible, and unevenly translated into frontier AI power. The dashboard reports 14.3 granted AI patents per 100,000 people in 2024, compared with 7.0 for China and 4.7 for the United States in the same comparison set. AI research publications rose from 14,348 in 2020 to 19,283 in 2025, while notable AI models increased sharply in 2025 after several uneven years. Patent density and publication volume indicate a real knowledge base, but they do not by themselves establish strategic capacity. Jeffrey Ding's diffusion argument suggests that national power follows when invention is absorbed across production systems, institutions, and users. For Korea, the research question is therefore not whether it can produce knowledge, but whether that knowledge changes manufacturing routines, defense validation, public administration, and platform services.",
        ["dashboard", "ding"],
    )
    add_figure(
        doc,
        FIG_DIR / "figure_5_research_models_trend.png",
        "Figure 4. Korea combines broad research volume with uneven but rising model output.",
    )
    add_para(
        doc,
        "The AI co scientist agenda is strategically important because it directs AI capability into fields where Korea already has industrial and scientific depth. MSIT's 2026 plan targets life sciences, earth sciences, mathematics, materials and chemistry, semiconductors and displays, and secondary batteries. The value of this agenda is not only discovery. It can generate datasets, laboratory routines, benchmarks, and domain specific methods that firms later absorb into production. In Bush's terms, Korea needs basic research as a long horizon national asset; in Ding's terms, that research becomes power only when it improves diffusion capacity.",
        ["work_plan", "bush", "ding"],
    )
    add_heading(doc, "Adoption Layer: Robots, Skills, and Physical Deployment", 2)
    add_para(
        doc,
        "Korea's adoption layer is the strongest argument for a distinctive AI strategy because the country can connect software intelligence to physical production faster than many competitors. The dashboard shows high robot density, rising AI engineering skills diffusion, broad worker experimentation with generative AI, and measurable AI diffusion across the economy. Yet the flat trend in annual robot installations suggests that adoption is reaching a more difficult stage. Korea must now integrate perception, planning, simulation, safety assurance, and human workflow redesign into already automated environments.",
        ["dashboard"],
    )
    add_figure(
        doc,
        FIG_DIR / "figure_6_adoption_skills_trend.png",
        "Figure 5. Korea's physical adoption base is deeper than its software skills layer.",
    )
    add_para(
        doc,
        "The 2026 Work Plan's emphasis on regional AI transformation and Physical AI fits Korea's industrial structure. MSIT's Physical AI Global Alliance, launched in September 2025 with participation from government, industry, universities, research institutes, and demand side firms such as Hyundai Motor, HD Hyundai Heavy Industries, and LG AI Research, shows that the state is trying to organize physical deployment as a national arena rather than a collection of firm projects. Manufacturing, logistics, shipbuilding, mobility, semiconductors, displays, and batteries provide demanding test environments where AI can raise productivity and produce exportable deployment templates. The state should treat these environments as national test beds for safety cases, interoperable tools, and sector benchmarks. If Korea succeeds, its comparative advantage will lie in reliable AI systems that operate in factories, vehicles, hospitals, ports, and defense networks.",
        ["work_plan", "physical_ai", "dashboard"],
    )
    add_heading(doc, "Startup and Funding Layer", 2)
    add_para(
        doc,
        "The startup and funding layer is broad enough to matter but not yet deep enough to carry national strategy without public demand creation. The dashboard records about 329 newly funded AI companies from 2013 through 2025 and total AI investment of about USD 10.75 billion over that period, with 2025 private AI investment at about USD 1.78 billion. The domain pattern is concentrated in health care and life sciences, business services, general purpose AI, and consumer services. Defense, public governance, industrial infrastructure, and education receive smaller amounts in the available investment data, even though they are central to the strategic logic of Korean AI.",
        ["dashboard"],
    )
    add_para(
        doc,
        "This funding pattern is not a market failure in the abstract; it reflects the revenue horizons of private capital. Strategic sectors often require regulated data, public procurement, liability rules, testing infrastructure, and patient reference customers before startups can scale. Korea's policy task is to convert public missions into credible markets without insulating weak firms from competition. Procurement sandboxes, milestone contracts, public data trusts, and access to national compute can link startups to defense, health, manufacturing, and regional transformation while preserving pressure to meet technical and commercial benchmarks.",
        ["ntsf", "dashboard"],
    )
    add_para(
        doc,
        "The political economy of Korean AI makes this design especially important because chaebol and platform firms are necessary for scale but can also become gatekeepers. Samsung, SK, Hyundai, LG, Naver, Kakao, and telecom groups can provide data centers, customer bases, devices, factories, and export channels that smaller firms cannot reproduce. The state should therefore use consortium support to create shared infrastructure while using procurement and interface rules to keep startups from becoming dependent subcontractors. Korea's older developmental bargain between the state and large firms needs a digital supplement that preserves contestability in models, tools, and applications.",
        ["ntsf", "models", "nvidia"],
    )
    add_figure(
        doc,
        FIG_DIR / "figure_3_korea_private_ai_focus.png",
        "Figure 6. Korea's private AI investment is concentrated in a few domains.",
    )
    add_heading(doc, "Strategic Assessment of Capacity", 2)
    add_para(
        doc,
        "Korea's capacity building strategy is strongest where it exploits the interaction of memory, manufacturing, trusted data governance, and applied deployment. It is weakest where it attempts to imitate superpower scale without controlling the full accelerator, cloud, and model ecosystem. A credible Korean strategy should therefore define success as specialized strategic indispensability rather than general AI supremacy. In practical terms, that means becoming the allied supplier of trusted memory and industrial AI deployment, the regional platform for regulated AI services, and a military innovator in sensor rich defense systems.",
        ["ntsf", "stanford", "dashboard"],
    )
    add_para(
        doc,
        "The NTSF pillars also reveal the main institutional risk. The state has begun to coordinate, science produces dense technical output, and industry has world class assets, but each pillar can still optimize for its own metrics. Ministries can count GPUs, researchers can count papers, and firms can count closed customers without producing national strategic power. Korea's next phase should therefore measure capacity by diffusion into strategically important uses: factory productivity, defense readiness, public service efficiency, startup access to compute, trusted data reuse, and leverage in allied technology bargaining.",
        ["ntsf", "ding"],
    )

    add_heading(doc, "Historical Lessons")
    add_para(
        doc,
        "The historical cases clarify when states should compel cooperation, absorb systemic risk, or build absorptive capacity rather than simply subsidize firms. Each case offers a different institutional mechanism for translating technology into power. The synthetic rubber program shows how to pool competitors around a scarce input, the Manhattan Project shows when the state must bear extraordinary risk, Meiji Japan shows how imported technology can be absorbed without permanent dependence, and Vannevar Bush shows why scientific capacity must be treated as a long horizon strategic asset.",
        ["rubber", "wellerstein", "ohno", "bush"],
    )
    add_heading(doc, "The United States Synthetic Rubber Program", 2)
    add_para(
        doc,
        "The synthetic rubber program demonstrates how a state can transform a supply shock into a temporary cooperation regime among competitors. The United States faced the loss of Southeast Asian natural rubber during the Second World War, then forced technical cooperation, pooled knowledge across firms, financed capacity, and guaranteed demand. The lesson for Korea is institutional rather than sectoral. A strategic technology program should identify the scarce input and compel cooperation where national need exceeds private incentives. In AI, the equivalent bottlenecks are shared access to compute, data, evaluation, memory integration, and sectoral deployment tools.",
        ["rubber", "ntsf"],
    )
    doc.add_page_break()
    add_heading(doc, "The Manhattan Project", 2)
    add_para(
        doc,
        "The Manhattan Project identifies the conditions under which state project authority becomes defensible: extraordinary strategic urgency, high systemic risk, and a technology whose development requires coordinated scientific, industrial, and military capacity. It connected science, industry, and military command under extreme secrecy and time pressure, while also creating a technology whose strategic effects outran political control. For Korean AI, the transferable lesson is not secrecy, but state responsibility for scale, testing, safety, and production infrastructure when markets cannot absorb the downside risk. Defense AI, model safety, and critical infrastructure AI belong in this category.",
        ["wellerstein", "ntsf"],
    )
    add_heading(doc, "Meiji Japan", 2)
    add_para(
        doc,
        "Meiji Japan matters for Korea because it separates technological dependence from strategic dependence. Japan used foreign technology under external pressure, but it built domestic absorptive capacity through education, institutional reform, industry, and adaptation to local conditions. Kenichi Ohno stresses that Japan's catch up combined broad learning, private dynamism, and policy support. Korea's AI strategy should follow the same logic by importing global tools where they accelerate learning, building domestic layers where dependence creates leverage for others, and adapting AI to Korean language, law, culture, manufacturing, and defense needs.",
        ["ohno"],
    )
    add_heading(doc, "Science and the Endless Frontier", 2)
    add_para(
        doc,
        "Vannevar Bush's science compact clarifies why AI strategy cannot be built only around visible applications and procurement targets. AI depends on long term scientific capacity, measurement infrastructure, talent formation, and discovery routines that markets may underprovide when the payoff is uncertain. Korea's AI co scientist agenda, AI safety work, and evaluation datasets should therefore be insulated from short political cycles. These investments create the knowledge base that industry later turns into products, standards, and operational competence.",
        ["bush", "work_plan"],
    )

    doc.add_page_break()
    add_heading(doc, "Comparative Country Approaches")
    add_para(
        doc,
        "Comparative cases show that middle power AI strategy works best when sovereignty is defined as selective control over strategic functions rather than isolation from global networks. Taiwan turns foundry centrality into geopolitical relevance. Singapore uses public digital infrastructure, trust, and global convening to shape standards and adoption. Israel links defense demand, venture capital, universities, and elite talent, while also exposing the risk of dependence on external capital and markets. Korea can draw from each model without copying any single national path because its own leverage lies in memory, manufacturing, platforms, and allied demand.",
        ["taiwan", "singapore", "israel"],
    )
    add_table(
        doc,
        ["Country or case", "Closest NTSF model", "Comparative insight for Korea"],
        [
            [
                "South Korea",
                "Hybrid of Model A, Model B, and Model D",
                "Korea is using consortium tools for K AI, national champion tools for chaebols, and commercial dual use tools for industrial AI. The missing step is a common mission layer that links these tools.",
            ],
            [
                "Taiwan",
                "Model B and Model D",
                "Taiwan turns foundry centrality into strategic relevance. Korea should do the same with memory, packaging, AI servers, and industrial AI deployment.",
            ],
            [
                "Singapore",
                "Model D with a state platform",
                "Singapore emphasizes trusted digital infrastructure, public good AI, global networks, and user adoption. Korea can use the AI Basic Act and K AI exports in a similar way.",
            ],
            [
                "Israel",
                "Model C and Model D",
                "Israel links military demand, technical talent, venture capital, and exportable security products. Korea can adapt the defense innovation part while building stronger civilian diffusion.",
            ],
            [
                "United States",
                "Model C and Model D",
                "The United States shows how public research and procurement can create platforms later scaled by industry. Korea needs similar demand creation for AI in public services and defense.",
            ],
        ],
        "Table 1. Comparative approaches and NTSF model fit.",
    )

    add_heading(doc, "Critique of Korea's Current Strategy")
    add_para(
        doc,
        "Korea's current strategy contains credible instruments, but their strategic value depends on whether they are coupled to missions that create learning, demand, and bargaining power. The weakness is not the absence of ambition or funding. It is the risk that model competitions, GPU purchases, regulatory reforms, and industrial programs proceed as separate policy tracks. The following critiques identify where institutional design must become more demanding.",
        ["ntsf", "budget", "models"],
    )
    add_numbered(
        doc,
        [
            (
                "The strategy remains vulnerable to AI G3 symbolism. Ranking third in notable model counts is useful, and patent density is meaningful, but statecraft requires assets that other countries need and cannot easily replace. Korea's stronger leverage lies in memory, manufacturing, deployment capability, and trusted data governance. Model prestige becomes strategically relevant only when it is connected to those control points.",
                ["stanford", "dashboard"],
            ),
            (
                "The model strategy risks fragmentation because five teams can duplicate infrastructure while producing incompatible application ecosystems. Competition can improve performance, but it should operate inside shared rules for interfaces, evaluation, sectoral benchmarks, safety reporting, and procurement compatibility. Otherwise public support may deepen market segmentation rather than create a national deployment layer.",
                ["models"],
            ),
            (
                "Compute strategy remains externally dependent in ways that procurement alone cannot solve. NVIDIA GPUs are necessary for near term capacity, but a sovereign strategy must also support domestic NPUs, Korean cloud operators, memory integration, energy planning, and reciprocal access to allied compute. GPU purchases create capacity; the surrounding ecosystem determines whether that capacity becomes leverage.",
                ["nvidia", "work_plan"],
            ),
            (
                "Data governance is both a constraint on firms and a possible source of institutional advantage. Korea has strong platforms and a serious privacy regulator. The danger is an asymmetric outcome in which domestic firms face slow data access while foreign models continue to reach Korean users. The strategic response is to make privacy preserving Korean data access fast enough for innovation and trusted enough for export.",
                ["pipc", "ai_basic"],
            ),
            (
                "Industrial adoption is the decisive test because Korea's advantage lies in diffusing AI into physical production and public services. Robot density and large manufacturers are important, but AI value will come from small suppliers, hospitals, schools, logistics, shipyards, battery firms, and defense units. Diffusion requires procurement, standards, training, and integration support, not only model access.",
                ["dashboard", "ding"],
            ),
            (
                "Defense AI can become a source of advantage only if Korea builds learning loops between civilian innovation and military validation. The Defense AI Center should create test ranges, data pipelines, realistic exercises, model assurance, and acquisition pathways for responsible adoption. Military value will depend on repeated operational evaluation rather than isolated prototypes.",
                ["defense", "horowitz"],
            ),
        ],
    )

    add_heading(doc, "Policy Recommendations")
    add_para(
        doc,
        "The policy agenda should move Korea from capability accumulation to institutional leverage. The state already has legal authority, public finance, industrial partners, and technology assets. Its next task is to organize these assets around missions that create bargaining power abroad and productivity gains at home. The recommendations below treat AI as a stack of interdependent inputs, institutions, and deployment channels rather than as a single model race.",
        ["ntsf", "ai_basic"],
    )
    add_numbered(
        doc,
        [
            (
                "Define Korea's AI strategy as stack sovereignty. The state should publish a stack map that covers memory, packaging, accelerators, cloud, data, models, applications, standards, talent, and export channels. Each layer should have a lead ministry, industrial partners, measurable bottlenecks, and a statecraft objective. This would make interdependence visible and prevent model development from crowding out less visible sources of leverage.",
                ["ntsf"],
            ),
            (
                "Use high bandwidth memory as alliance currency. Seoul should negotiate supply assurance deals with the United States, Japan, Taiwan, and trusted European partners. In exchange for dependable memory access, Korea should seek reciprocal compute access, co investment in domestic data centers, favorable treatment for Korean AI cloud providers, and allied demand for Korean AI semiconductors. Memory exports should become part of a broader bargain over AI infrastructure, market access, and allied resilience.",
                ["trendforce", "nvidia"],
            ),
            (
                "Create a National AI Compute and Data Commons. The commons should allocate public GPUs and selected private cloud capacity to startups, universities, public agencies, and strategic sector projects. Access should be tied to Korean evaluation benchmarks, privacy rules, security standards, and shared learning obligations. The synthetic rubber lesson applies because scarce inputs should be pooled where duplication wastes national capacity, while firms still compete in applications and services.",
                ["rubber", "budget"],
            ),
            (
                "Turn the K AI model competition into a shared application platform. The five model teams should compete on performance inside common rules for APIs, evaluation formats, safety reporting, Korean language and culture tests, domain adapters, and public procurement compatibility. Government demand should reward working services in public agencies and strategic sectors, which would discipline the model program through use rather than publicity.",
                ["models"],
            ),
            (
                "Make manufacturing the first national AI mission. Korea should build AI deployment programs for semiconductors, displays, batteries, shipbuilding, autos, steel, logistics, and robotics. Each mission should include large firms, suppliers, universities, system integrators, and regional governments. The goal should be productivity diffusion across the supply base, because national advantage depends on raising the capability of suppliers and regions as well as the largest chaebols.",
                ["dashboard", "work_plan"],
            ),
            (
                "Use defense as a disciplined adopter. The Defense AI Center should run joint test ranges with industry and academia for surveillance, logistics, cyber defense, drone coordination, maintenance, decision support, and training. Systems should pass red team testing, human control review, data provenance checks, and operational exercises before procurement. This would make defense demand a source of validation for civilian technology rather than a closed procurement track.",
                ["defense", "horowitz"],
            ),
            (
                "Build trusted data spaces for health, mobility, manufacturing, finance, and public services. The Personal Information Protection Commission should create safe harbor templates for privacy preserving research, federated learning, synthetic data, and secure evaluation. These institutions would let firms use sensitive data under clear rules while reducing legal uncertainty and creating a Korean governance export.",
                ["pipa", "pipc", "ai_basic"],
            ),
            (
                "Make the AI Basic Act a statecraft product. Korea should export implementation guidance, transparency standards, labeling methods, high impact AI review practices, and Korean language safety benchmarks to partners in Southeast Asia and the broader Indo Pacific. Regulation can become influence when other states adopt Korean methods for trusted deployment, similar to Singapore's use of digital governance as diplomatic capital.",
                ["ai_basic", "singapore"],
            ),
            (
                "Create a national talent absorption program. Korea should expand AI graduate fellowships, allow defense related AI service pathways, recruit overseas Korean researchers, and make immigration easier for AI engineers, chip designers, data center operators, and AI safety researchers. Talent policy should address the broad engineering workforce required for diffusion as well as the smaller group of frontier researchers.",
                ["oecd", "ntsf"],
            ),
            (
                "Put energy and grid capacity inside AI policy. Data centers, AI factories, and supercomputing need reliable electricity, cooling, land, and permitting. Seoul should link AI clusters to regional power planning, nuclear and renewable capacity, grid upgrades, and water management. Compute policy that ignores power and permitting will create nominal capacity without operational sovereignty.",
                ["work_plan", "taiwan"],
            ),
            (
                "Build an AI assurance industry. Korea should fund firms and labs that test models for security, privacy, bias, reliability, Korean language performance, physical safety, and military use. Assurance can become an exportable service for countries that want trusted AI and prefer diversified evaluation providers.",
                ["ai_basic", "nscai"],
            ),
            (
                "Shift from subsidy logic to procurement logic. Grants are useful at the invention stage, but industrial diffusion needs buyers. The government should publish multi year procurement road maps for AI in hospitals, disaster response, tax administration, customs, defense logistics, education support, and industrial safety. Demand clarity would let firms invest in sectoral systems rather than generic demonstrations.",
                ["budget", "bush"],
            ),
            (
                "Protect startups from chaebol capture. National champions are necessary for scale, but Korea should reserve compute credits, public datasets, procurement lots, and export support for startups and mid sized firms. The state should require large consortiums to publish interfaces and partner terms so smaller firms can build on K AI rather than become dependent subcontractors.",
                ["models", "israel"],
            ),
            (
                "Create a Korean AI statecraft unit. This unit should sit above single ministries and track where Korea has leverage, vulnerability, and bargaining power. It should monitor memory dependence, export controls, cloud dependence, talent flows, standards votes, foreign model penetration, and defense readiness. Quarterly reporting to the National AI Strategy Committee would force technology policy to be assessed as a statecraft portfolio rather than a set of ministry programs.",
                ["ntsf", "ai_committee", "ai_basic"],
            ),
        ],
    )

    add_heading(doc, "Conclusion")
    add_para(
        doc,
        "Korea's strongest path to AI power is selective, infrastructural, and adoption centered. The country can supply critical memory, deploy AI deeply into advanced manufacturing, build trusted Korean and sectoral models, adapt defense to demographic scarcity, and export a practical governance model for allied and middle power states. This path does not require Korea to imitate the scale strategies of the United States or China. It requires Korea to turn its narrower set of control points into bargaining power, productivity gains, and trusted deployment capacity.",
        ["stanford", "dashboard"],
    )
    add_para(
        doc,
        "The NTSF framework points to integration as the central policy problem. Technology strategy succeeds when the state defines the problem, science expands the frontier, industry scales the solution, and the whole system translates capacity into power. Korea has each pillar, but the pillars must be organized around missions that create feedback among compute access, privacy preserving data, foundation models, manufacturing deployment, defense adoption, memory leverage, and standards diplomacy. AI will become a central instrument of Korean statecraft only if those feedback loops are deliberately built.",
        ["ntsf"],
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "Appendix")
    add_para(
        doc,
        "The NTSF models clarify that the policy choice is not whether the state or the market should lead, but which coordination form fits the strategic constraint being addressed. Korea's AI strategy requires different models for different layers of the stack because compute scarcity, chaebol scale, defense adoption, and commercial diffusion do not respond to the same instrument.",
        ["ntsf"],
    )
    add_table(
        doc,
        ["NTSF model", "Strategic context", "Mechanism", "Korea AI implication"],
        [
            [
                "Model A: Consortium",
                "National emergency or scarce strategic input",
                "The state pools competitors, funds capacity, and guarantees demand.",
                "Use for compute, data, evaluation, and strategic AI tools where duplication wastes scarce resources.",
            ],
            [
                "Model B: National Champion",
                "Catch up modernization and scale competition",
                "The state supports large firms in exchange for national objectives.",
                "Use Samsung, SK, Hyundai, LG, Naver, and SK Telecom, while preventing closed ecosystems from blocking startups.",
            ],
            [
                "Model C: State Project",
                "Secret, military, or public good technology",
                "The state is the main buyer and absorbs systemic risk.",
                "Use for defense AI, AI safety infrastructure, public service AI, and critical infrastructure assurance.",
            ],
            [
                "Model D: Commercial Dual Use",
                "Commercial technology with strategic value",
                "The state shapes markets through procurement, standards, export rules, and security compensation.",
                "Use for industrial AI, cloud, platforms, health, finance, mobility, robotics, and AI governance exports.",
            ],
        ],
        "Table 2. The four NTSF models for technology strategy.",
    )

    add_heading(doc, "Notes")
    add_para(
        doc,
        "The notes record the official policy sources, course materials, data sources, and comparative references used to support the factual and analytical claims in the paper.",
    )
    for idx, key in enumerate(ORDERED_NOTE_KEYS, start=1):
        note = NOTES[key]
        reject_dashes(note, "note")
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{idx}. ")
        r.bold = True
        p.add_run(note)

    add_heading(doc, "Selected Bibliography")
    add_para(
        doc,
        "The bibliography prioritizes official Korean policy materials, class frameworks, international datasets, and comparative technology strategy sources rather than general commentary on AI.",
    )
    bibliography = [
        NOTES["rubber"],
        NOTES["strategy_2019"],
        NOTES["digital_strategy"],
        NOTES["ai_basic"],
        NOTES["ai_committee"],
        NOTES["budget"],
        NOTES["gpu_procurement"],
        NOTES["models"],
        NOTES["physical_ai"],
        NOTES["work_plan"],
        NOTES["oecd"],
        NOTES["stanford"],
        NOTES["trendforce"],
        NOTES["nvidia"],
        NOTES["pipa"],
        NOTES["horowitz"],
        NOTES["ding"],
        NOTES["ohno"],
        NOTES["wellerstein"],
        NOTES["bush"],
        NOTES["nscai"],
        NOTES["singapore"],
        NOTES["taiwan"],
        NOTES["israel"],
        NOTES["dashboard"],
    ]
    for item in bibliography:
        reject_dashes(item, "bibliography")
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.left_indent = Inches(0.25)
        p.add_run(item)

    return doc


def audit_docx_text(path: Path) -> list[str]:
    from docx import Document as ReadDocument

    doc = ReadDocument(path)
    texts: list[str] = []
    for p in doc.paragraphs:
        texts.append(p.text)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                texts.append(cell.text)
    return [x for x in texts if re.search(r"[-–—]", x)]


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for note in NOTES.values():
        reject_dashes(note, "note")
    doc = build_doc()
    output_path = DOCX_PATH
    try:
        doc.save(output_path)
    except PermissionError:
        output_path = OUT_DIR / "korea_ai_strategy_paper_revised.docx"
        doc.save(output_path)
    offenders = audit_docx_text(output_path)
    if offenders:
        raise ValueError("Generated document contains dash characters: " + repr(offenders[:10]))
    print(output_path)


if __name__ == "__main__":
    main()
