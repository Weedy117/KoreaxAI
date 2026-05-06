from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

from build_korea_ai_strategy_paper import (
    NOTES,
    ORDERED_NOTE_KEYS,
    add_citations,
    add_figure,
    add_heading,
    add_para,
    audit_docx_text,
    reject_dashes,
    style_document,
)


OUT_DIR = Path(r"C:\Users\ASUS\Desktop\SAIS Playground\ROKxAI\outputs\korea_ai_strategy")
FIG_DIR = OUT_DIR / "figures"
DOCX_PATH = OUT_DIR / "korea_ai_strategy_paper.docx"
FALLBACK_DOCX_PATH = OUT_DIR / "korea_ai_strategy_paper_consolidated.docx"


NOTES.update(
    {
        "action_plan": "National Artificial Intelligence Strategy Committee and Ministry of Science and ICT, Republic of Korea Artificial Intelligence Action Plan, February 2026.",
        "carnegie": "Darcie Draudt Vejares and Seungjoo Lee, Governing AI in the Shadow of Giants: Korea's Strategic Response to Great Power AI Competition, Carnegie Endowment for International Peace, April 30, 2026.",
        "stimson": "Sean Hyuk Hyun Kwon and J. James Kim, From Compute to Capacity: South Korea's Approach to Industrial AI Adoption, Stimson Center, February 6, 2026.",
        "stimson_manufacturing": "Sean Hyuk Hyun Kwon, Revolutionizing the Industrial Base: South Korea's AI Integration in Manufacturing, Stimson Center, February 9, 2026.",
        "field": "Alexander J. Field, The World War Two U.S. Rubber Famine, forthcoming in Business History Review, 2026.",
        "yonhap_models": "Yonhap News Agency, report on Korea's sovereign AI foundation model competition, January 15, 2026.",
        "ifrs": "International Federation of Robotics, World Robotics public data and press materials on robot density, 2025.",
        "action_news": "Kyunghyang Shinmun, report on final approval of the Republic of Korea Artificial Intelligence Action Plan and K Moonshot, February 25, 2026.",
        "copyright_debate": "Dong A Ilbo, report on Korea Newspaper Association objections to AI Action Plan copyright proposals, January 5, 2026.",
    }
)


def add_note_section(doc: Document) -> None:
    add_heading(doc, "Notes")
    for idx, key in enumerate(ORDERED_NOTE_KEYS, start=1):
        note = NOTES[key]
        reject_dashes(note, "note")
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{idx}. ")
        r.bold = True
        p.add_run(note)


def add_compact_bibliography(doc: Document) -> None:
    add_heading(doc, "Selected Bibliography")
    bibliography = [
        NOTES["action_plan"],
        NOTES["ai_basic"],
        NOTES["carnegie"],
        NOTES["stimson"],
        NOTES["rubber"],
        NOTES["field"],
        NOTES["wellerstein"],
        NOTES["bush"],
        NOTES["ohno"],
        NOTES["ding"],
        NOTES["horowitz"],
        NOTES["singapore"],
        NOTES["israel"],
        NOTES["stanford"],
        NOTES["oecd"],
        NOTES["dashboard"],
    ]
    for item in bibliography:
        reject_dashes(item, "bibliography")
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def build_doc() -> Document:
    ORDERED_NOTE_KEYS.clear()
    doc = Document()
    style_document(doc)

    title = "Korea's AI Strategy as Statecraft"
    subtitle = "Selective Leverage in Memory, Manufacturing, Data, and Defense"
    reject_dashes(title, "title")
    reject_dashes(subtitle, "subtitle")
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    q = doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    qr = q.add_run(subtitle)
    qr.font.size = Pt(14)
    qr.font.color.rgb = RGBColor(80, 80, 80)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Policy research paper for Technology, Power, and Statecraft").italic = True
    doc.add_paragraph()

    add_heading(doc, "Executive Argument")
    add_para(
        doc,
        "South Korea's AI strategy is best understood as selective stack statecraft rather than a bid for comprehensive technological autonomy. The government's aspiration to become a top three AI power is politically useful, but it is a poor analytical benchmark if read as model parity with the United States or China. Korea's more plausible route to AI power is to make itself difficult to substitute in the layers where it already has leverage: high bandwidth memory and advanced packaging, precision manufacturing, robotics and physical AI, trusted Korean data governance, defense adaptation, and public sector deployment. AI becomes statecraft when these assets produce bargaining leverage abroad, productivity gains under demographic pressure at home, and military adaptation against a proximate threat.",
        ["stanford", "dashboard", "carnegie"],
    )
    add_para(
        doc,
        "This reading narrows both ambition and risk. Korea has an unusually dense innovation base, leading AI patents per capita in the Stanford AI Index dataset used for this project, and it has world leading industrial robot density. Yet foundation model production, hyperscale cloud, frontier compute, and data center energy remain concentrated elsewhere. The strategic question is therefore not whether Korea can own the whole AI stack. It is whether Seoul can coordinate the stack layers it controls well enough to shape the terms on which Korean firms, allies, and public institutions use AI.",
        ["stanford", "dashboard", "ifrs"],
    )
    add_figure(
        doc,
        FIG_DIR / "figure_1_patents_robot_density.png",
        "Figure 1. Korea's measurable AI advantages are concentrated in innovation density and industrial automation.",
    )

    doc.add_page_break()
    add_heading(doc, "Strategic Landscape")
    add_para(
        doc,
        "The policy landscape shows that Seoul now imagines AI less as a discrete software sector than as an execution architecture for national renewal. The 2019 National Strategy for Artificial Intelligence framed AI through ecosystem creation, utilization, and people centered policy. The Data Dam under the Digital New Deal treated data and AI as infrastructure for recovery and digital modernization. The 2022 National Digital Strategy widened the frame to chips, data, platforms, government, and talent. The AI Basic Act and the 2026 Republic of Korea Artificial Intelligence Action Plan move beyond that sequence by turning AI into a whole of government implementation problem. The Action Plan's three pillars, twelve strategic areas, ninety nine implementation tasks, and three hundred twenty six policy recommendations reveal a state trying to convert ambition into assigned ministries, deadlines, legal changes, and administrative roles.",
        ["strategy_2019", "digital_strategy", "ai_basic", "action_plan"],
    )
    add_para(
        doc,
        "The inference is that Korea is reviving a developmental state habit under digital conditions, but with a different coordination problem. Earlier industrial policy could concentrate credit, imports, and export discipline around identifiable sectors. AI requires coordination across compute, data, cloud, semiconductors, copyright, privacy, safety, public procurement, schools, defense, and regional industry. The Presidential Council on National Artificial Intelligence Strategy, Chief AI Officer structures, expert committees, and the AI Basic Act's enforcement machinery are instruments for making those dependencies administratively visible. Their weakness is also visible: a plan with nearly one hundred tasks can become a catalogue of ministry activity unless the control tower ranks bottlenecks and forces tradeoffs.",
        ["ai_basic", "action_plan", "carnegie"],
    )
    add_para(
        doc,
        "Korea's geopolitical position reinforces this selective logic. Carnegie's 2026 assessment describes Korea as pursuing a dual full stack strategy: domestic capability building across key AI layers while embedding itself in a United States centered ecosystem. That formulation captures the constraint. Korea cannot detach from American GPUs, cloud platforms, and frontier model ecosystems, but it can use memory, manufacturing, robotics, and allied production capacity to gain leverage within them. NVIDIA's announced Korea infrastructure commitments and the role of Samsung Electronics and SK hynix in advanced memory make Korea a supplier and user of AI infrastructure rather than a simple rule taker. The same dependence creates vulnerability because announced GPUs are not the same as sovereign compute, and memory strength does not eliminate reliance on foreign accelerators, software stacks, and data center capacity.",
        ["carnegie", "nvidia", "trendforce"],
    )
    add_para(
        doc,
        "Domestic pressures explain why the Action Plan emphasizes AI transformation rather than model prestige alone. Korea's working age population is projected to contract sharply over the coming decades, and the OECD and Korea Labor Institute find that small and medium enterprise AI adoption remains uneven despite Korea's large industrial base. This makes AI a productivity and diffusion problem before it is a model race. Physical AI, manufacturing AI, welfare automation, public service AI, and defense AI appear repeatedly in recent plans because they promise to convert scarce labor, high skill manufacturing, and accumulated process data into national capacity. The policy risk is that large firms move first while suppliers, regions, schools, hospitals, and defense acquisition systems lag behind.",
        ["oecd", "stimson", "action_plan"],
    )

    doc.add_page_break()
    add_heading(doc, "Capacity Building")
    add_para(
        doc,
        "Korea's capacity building strategy is strongest where it joins law, industrial assets, and deployment demand; it is weakest where it treats capability accumulation as a substitute for institutional integration. The AI Basic Act, now in force, gives the state a framework for promotion, transparency, high impact AI, safety obligations for highly capable systems, testing infrastructure, data support, overseas expansion, and national AI governance. Its minimum regulation principle shows that Seoul defines trustworthiness as a condition of industrial scaling rather than as an external restraint on the AI sector. The law remains a framework statute, however. Its strategic value depends on enforcement decrees, grace period guidance, PIPC practice, procurement rules, and whether firms can understand compliance without slowing legitimate deployment.",
        ["ai_basic", "pipa"],
    )
    add_para(
        doc,
        "The compute and model layer reveals the clearest attempt to build sovereign capacity and the clearest evidence of dependency. The Action Plan and MSIT work plan call for large public GPU allocations, AI Highway infrastructure, domestic AI semiconductors, and Korean foundation models. MSIT initially selected five sovereign foundation model teams, but early 2026 reporting showed the competition narrowing after Naver Cloud and NC AI failed the first evaluation round, with SK Telecom, LG AI Research, and Upstage advancing. The episode is analytically useful because it exposes the limits of sovereignty defined as domestic model branding. A model race can discipline performance, originality, and open release commitments, but sovereignty is hollow if the resulting systems cannot access lawful training data, domestic or allied compute, Korean evaluation benchmarks, and sector deployment channels.",
        ["action_plan", "models", "work_plan", "yonhap_models"],
    )
    add_para(
        doc,
        "Data governance is being recast as infrastructure rather than as a privacy obstacle. The Action Plan links national data platforms, AI training data, health data, public records, copyright reform, security disclosure, and AI based welfare. The PIPC's review of Kakao's Kanana service illustrates the Korean method: permit AI service development under ex ante safeguards, deidentification or encryption where appropriate, proprietary model preference for sensitive uses, and output controls. The unresolved tension is copyright and public data reuse. Creator groups and civil society have criticized proposals that would reduce legal uncertainty for AI training without sufficiently protecting rights. Korea's advantage will come from making Korean data reusable under credible rules, not from treating data protection as an administrative hurdle to be cleared.",
        ["action_plan", "pipc", "pipa", "copyright_debate"],
    )
    add_para(
        doc,
        "Physical AI is the most Korea specific part of the strategy because it aligns with industrial structure rather than imitating Silicon Valley. MSIT's Physical AI Global Alliance and the Action Plan's goal of global leadership in physical AI by 2030 connect robotics, vehicles, AI semiconductors, manufacturing, wellness technology, talent, governance, and computing resources. Stimson's analysis reaches a similar conclusion from the expert side: Korea's AI future will depend less on who builds the largest general model and more on whether AI can create scalable, trusted manufacturing capacity. Shipbuilding, defense manufacturing, automotive production, semiconductor fabrication, batteries, and logistics are not examples added after the model strategy. They are the places where Korea can turn deployment into data, quality control, supplier discipline, and alliance value.",
        ["physical_ai", "stimson", "stimson_manufacturing"],
    )
    add_para(
        doc,
        "Innovation, startups, and talent complete the picture but also expose the distributional weakness of a chaebol centered system. Korea's patent density and platform firms show real technical capacity, and the sovereign model project includes startups such as Upstage alongside conglomerate groups and research universities. Yet AI diffusion will depend on whether small suppliers, regional manufacturers, public hospitals, schools, and defense units can buy and adapt working systems. Singapore's digital government model and Israel's deep tech finance show different strengths: Singapore converts trusted digital infrastructure into adoption and diplomatic credibility, while Israel links venture capital, defense demand, and research talent. Korea cannot simply copy either model because its comparative advantage lies in production scale and chaebol supply chains. It can, however, use procurement, testing infrastructure, compute credits, standards, and public data access to prevent domestic AI from becoming only a large firm strategy.",
        ["dashboard", "singapore", "israel", "stimson"],
    )
    add_figure(
        doc,
        FIG_DIR / "figure_6_adoption_skills_trend.png",
        "Figure 2. Korea's adoption base is physical, while AI engineering skills are still diffusing.",
    )

    add_heading(doc, "Historical and Comparative Logic")
    add_para(
        doc,
        "The course cases are useful only if they identify mechanisms rather than analogies of scale. The United States synthetic rubber program shows how a state can respond when a strategic input becomes scarce and private competition alone cannot solve the bottleneck. After Japan cut off most natural rubber access, the United States used the Rubber Reserve Company, public plant construction, priority allocation, patent sharing, a common GRS formula, and demand assurance to compel cooperation among competing firms. The lesson for Korea is selective pooling. Compute, high quality Korean data, safety evaluation, and sector benchmarks should be treated like scarce strategic inputs where duplication wastes national capacity, while firms still compete downstream in models, applications, and services. The limit is equally important: rubber was a discrete wartime substitute target, whereas AI is a general purpose system with privacy, copyright, and competition concerns.",
        ["rubber", "field"],
    )
    add_para(
        doc,
        "The Manhattan Project offers a narrower lesson than crash mobilization rhetoric usually suggests. Wellerstein's account emphasizes that the project was not just science. Oak Ridge, Hanford, Los Alamos, universities, industrial contractors, and military administration converted uncertain research into production through systems integration under extraordinary risk. Korea should not imitate secrecy or total war governance. The relevant parallel is the organization of hard infrastructure around bottlenecks. Defense AI, model assurance, critical infrastructure AI, and AI for science require test ranges, evaluation institutions, procurement channels, and data pipelines that absorb risk no single firm can rationally carry alone.",
        ["wellerstein", "horowitz"],
    )
    add_para(
        doc,
        "Bush's Science, The Endless Frontier points to the opposite horizon: durable scientific capital. Bush argued that government could promote industrial research by expanding basic knowledge and scientific talent, not by managing every application. Korea's K Moonshot, AI co scientist agenda, semiconductor and materials AI, Korean language benchmarks, and safety research should be interpreted in that tradition. The caution is that Bush's university centered model is incomplete for AI, where capability also comes from deployed systems, open source communities, private labs, and operational data. Korea therefore needs a research base joined to diffusion institutions rather than research funding detached from industrial use.",
        ["bush", "work_plan", "action_news"],
    )
    add_para(
        doc,
        "Meiji Japan adds a learning mechanism: progressive absorption rather than imitation. Ohno's account describes a shift from foreign dependent turnkey projects toward Japanese owned operation through technical education, imported machinery, foreign advisers, domestic engineering institutions, and adaptation after failed copying. Korea is not a late industrializer in the Meiji sense, but the mechanism applies to AI dependence. Seoul can import GPUs, cloud services, open models, and foreign partnerships while deliberately localizing the layers where dependence would reduce leverage: Korean data spaces, high bandwidth memory integration, manufacturing AI, defense assurance, and Korean language systems. Comparative cases reinforce this point. Taiwan turns foundry centrality into strategic indispensability; Singapore turns trusted digital infrastructure into governance influence; Israel turns defense demand and venture finance into deep tech formation. Korea's distinctive model must combine all three logics without confusing selective sovereignty with autarky.",
        ["ohno", "taiwan", "singapore", "israel"],
    )

    add_heading(doc, "Conclusion")
    add_para(
        doc,
        "Korea's AI strategy has evolved from digital catch up policy into an attempt to govern a national AI stack. Its strongest path is selective, infrastructural, and adoption centered. It does not require Korea to match the United States or China in frontier model scale. It requires Korea to use its narrower control points to become valuable to allies, productive under demographic constraint, and institutionally credible in trusted deployment.",
        ["ding", "carnegie"],
    )
    add_para(
        doc,
        "Success should therefore be measured less by the slogan of AI top three than by concrete indicators of strategic conversion: lawful Korean data reuse, startup and university access to compute, SME participation in manufacturing AI, defense systems validated through realistic testing, productivity gains in shipyards and fabs, credible high impact AI oversight, and allied reliance on Korean memory and production capacity. Korea has the assets to make AI an instrument of statecraft, but only if the control tower turns a broad action plan into a ranked set of bottlenecks and missions.",
        ["action_plan", "oecd", "stimson"],
    )

    add_note_section(doc)
    add_compact_bibliography(doc)
    return doc


def word_count(path: Path) -> int:
    doc = Document(path)
    words: list[str] = []
    for p in doc.paragraphs:
        words.extend(re.findall(r"\b[\w']+\b", p.text))
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                words.extend(re.findall(r"\b[\w']+\b", cell.text))
    return len(words)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for note in NOTES.values():
        reject_dashes(note, "note")
    doc = build_doc()
    output_path = DOCX_PATH
    try:
        doc.save(output_path)
    except PermissionError:
        output_path = FALLBACK_DOCX_PATH
        doc.save(output_path)
    offenders = audit_docx_text(output_path)
    if offenders:
        raise ValueError("Generated document contains dash characters: " + repr(offenders[:10]))
    words = word_count(output_path)
    if words >= 4000:
        raise ValueError(f"Document word count exceeds target: {words}")
    print(f"{output_path} words={words}")


if __name__ == "__main__":
    main()
